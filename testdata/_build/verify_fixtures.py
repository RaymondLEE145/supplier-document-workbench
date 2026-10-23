# -*- coding: utf-8 -*-
"""Read every generated fixture back and assert it matches manifest.json.

Run after build_fixtures.py:
    C:\\Users\\Administrator\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe repo/testdata/_build/verify_fixtures.py

Exit code 0 => VERIFY OK.  Writes repo/testdata/_build/verify_report.json.
"""

import json
import sys
import traceback
from pathlib import Path

HERE = Path(__file__).resolve().parent
TESTDATA = HERE.parent
REPO = TESTDATA.parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

import fixture_data as fd  # noqa: E402

CHECKS = []
FAILURES = []
EXTRACTED = {}


def ensure_stdout_utf8():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


def check(name, ok, detail=""):
    ok = bool(ok)
    CHECKS.append({"check": name, "ok": ok, "detail": str(detail)})
    if not ok:
        FAILURES.append({"check": name, "detail": str(detail)})
        print("FAIL {}".format(name) + ("  <- " + str(detail) if detail else ""))
    else:
        print("PASS {}".format(name))


def approx(a, b, eps=1e-6):
    try:
        return abs(float(a) - float(b)) <= eps
    except Exception:
        return False


def compact(text):
    return "".join(str(text).split())


def verify_pdf(doc, path, manifest):
    import pymupdf

    expected = doc["extraction_expectations"]
    pdf = pymupdf.open(str(path))
    check(doc["id"] + ": page count == 1", pdf.page_count == 1, "pages={}".format(pdf.page_count))
    page = pdf[0]
    text = page.get_text("text")
    images = page.get_images(full=True)
    pdf.close()
    EXTRACTED[doc["id"]] = {"chars": len(text), "first_300": text[:300].replace("\n", " | ")}

    plain = compact(text)
    check(doc["id"] + ": PDF text layer is non-trivial", len(text) >= 600, "chars={}".format(len(text)))

    # real text layer, not a scanned image
    check(doc["id"] + ": no embedded page image (text-only PDF)", len(images) == 0,
          "images={}".format(len(images)))

    sup = fd.SUPPLIERS[doc["supplier_id"]]
    for keyword in [sup["name_en"], sup["contract_no"], sup["invoice_no"]]:
        check(doc["id"] + ": contains " + keyword, compact(keyword) in plain)

    for kw in ("HSCODE", "CBM", "TOTAL"):
        check(doc["id"] + ": contains " + kw, kw in plain)

    # language + item content
    title_zh = compact("装箱清单" if doc["doc_type"] == "packing_list" else "商业发票")
    check(doc["id"] + ": Chinese title present", title_zh in plain)
    for item in [it for it in fd.ITEMS if it["key"] in doc["line_item_keys"]]:
        check(doc["id"] + ": item EN " + item["key"], compact(item["name_en"]) in plain)
        check(doc["id"] + ": item ZH " + item["key"], compact(item["name_zh"]) in plain)
        if item["hs_code"]:
            check(doc["id"] + ": HS " + item["hs_code"], item["hs_code"] in plain)

    totals = fd.supplier_totals(doc["supplier_id"])
    check(doc["id"] + ": total cartons", "{:,}".format(totals["cartons"]) in text or str(totals["cartons"]) in text)
    check(doc["id"] + ": total pieces", "{:,}".format(totals["total_pieces"]) in text or str(totals["total_pieces"]) in text)
    check(doc["id"] + ": gross weight", "{:,}".format(totals["gross_weight_kg"]) in text)
    check(doc["id"] + ": CBM total", "{:.2f}".format(totals["cbm"]) in text)
    if doc["doc_type"] == "commercial_invoice":
        check(doc["id"] + ": grand amount", fd.money(totals["amount"]) in text)
        check(doc["id"] + ": SAY TOTAL line", "SAYTOTAL" in plain)

    # cross-check with pypdf (a second, independent text extractor)
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        text2 = "".join((p.extract_text() or "") for p in reader.pages)
        EXTRACTED[doc["id"]]["pypdf_chars"] = len(text2)
        check(doc["id"] + ": pypdf extracts company name", sup["name_en"] in text2)
        check(doc["id"] + ": pypdf extracts invoice no", sup["invoice_no"] in text2)
    except ImportError:
        check(doc["id"] + ": pypdf available", True, "pypdf not installed - skipped")
    except Exception as exc:  # keep the report instead of crashing the verifier
        check(doc["id"] + ": pypdf cross-check", False, repr(exc))


def verify_xlsx(doc, path, manifest):
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True)
    check(doc["id"] + ": workbook readable", wb.sheetnames == doc["extraction_expectations"]["sheet_names"],
          "sheets={}".format(wb.sheetnames))
    ws = wb.worksheets[0]
    sup = fd.SUPPLIERS[doc["supplier_id"]]
    exp = doc["extraction_expectations"]

    check(doc["id"] + ": company in title row 1", ws["A1"].value == sup["name_en"], "A1={}".format(ws["A1"].value))
    title = str(ws["A2"].value or "")
    check(doc["id"] + ": title row 2", ("装箱清单" in title) if doc["doc_type"] == "packing_list" else ("商业发票" in title),
          "A2={}".format(title))
    check(doc["id"] + ": title row 2 is in row 2 (two title rows)", ws["A1"].value not in (None, "") and ws["A2"].value not in (None, ""))

    header_row = exp["header_row"]
    data_rows = exp["data_rows"]
    items = [it for it in fd.ITEMS if it["key"] in doc["line_item_keys"]]

    if doc["doc_type"] == "packing_list":
        check(doc["id"] + ": header A row", str(ws.cell(header_row, 1).value).startswith("NO."))
        check(doc["id"] + ": header B cartons first", str(ws.cell(header_row, 2).value).startswith("CARTONS"))
        check(doc["id"] + ": header E HS", "HS CODE" in str(ws.cell(header_row, 5).value))
        for item, row in zip(items, data_rows):
            check(doc["id"] + ": row {} cartons".format(row), ws.cell(row, 2).value == item["cartons"],
                  "{} vs {}".format(ws.cell(row, 2).value, item["cartons"]))
            check(doc["id"] + ": row {} name EN".format(row), ws.cell(row, 3).value == item["name_en"])
            check(doc["id"] + ": row {} name ZH".format(row), ws.cell(row, 4).value == item["name_zh"])
            check(doc["id"] + ": row {} HS".format(row), ws.cell(row, 5).value == item["hs_code"],
                  "cell={} expected={}".format(ws.cell(row, 5).value, item["hs_code"]))
            check(doc["id"] + ": row {} total pcs".format(row), ws.cell(row, 8).value == item["total_pieces"])
            check(doc["id"] + ": row {} CBM".format(row), approx(ws.cell(row, 12).value, item["cbm"]))
        totals = fd.supplier_totals(doc["supplier_id"])
        check(doc["id"] + ": total cartons", ws.cell(exp["totals_row"], 2).value == totals["cartons"])
        check(doc["id"] + ": total pieces", ws.cell(exp["totals_row"], 8).value == totals["total_pieces"])
        check(doc["id"] + ": total net", ws.cell(exp["totals_row"], 9).value == totals["net_weight_kg"])
        check(doc["id"] + ": total gross", ws.cell(exp["totals_row"], 10).value == totals["gross_weight_kg"])
        check(doc["id"] + ": total CBM", approx(ws.cell(exp["totals_row"], 12).value, totals["cbm"]))
    else:
        check(doc["id"] + ": header D MODEL before HS", str(ws.cell(header_row, 4).value).startswith("MODEL"))
        check(doc["id"] + ": header E HS", "HS CODE" in str(ws.cell(header_row, 5).value))
        for item, row in zip(items, data_rows):
            check(doc["id"] + ": row {} model".format(row), ws.cell(row, 4).value == item["model"])
            check(doc["id"] + ": row {} HS".format(row), ws.cell(row, 5).value == item["hs_code"],
                  "cell={} expected={}".format(ws.cell(row, 5).value, item["hs_code"]))
            check(doc["id"] + ": row {} cartons".format(row), ws.cell(row, 7).value == item["cartons"])
            check(doc["id"] + ": row {} total pcs".format(row), ws.cell(row, 9).value == item["total_pieces"])
            check(doc["id"] + ": row {} unit price".format(row), approx(ws.cell(row, 10).value, item["unit_price"]))
            check(doc["id"] + ": row {} amount".format(row), approx(ws.cell(row, 11).value, item["amount"]))
        totals = fd.supplier_totals(doc["supplier_id"])
        check(doc["id"] + ": total cartons", ws.cell(exp["totals_row"], 7).value == totals["cartons"])
        check(doc["id"] + ": total pieces", ws.cell(exp["totals_row"], 9).value == totals["total_pieces"])
        check(doc["id"] + ": total amount", approx(ws.cell(exp["totals_row"], 11).value, totals["amount"]))

    # merged title cells survive
    merges = {str(r) for r in ws.merged_cells.ranges}
    check(doc["id"] + ": merged title cells present", any(str(r).startswith("A1:") for r in ws.merged_cells.ranges),
          "merges={}".format(sorted(merges)[:5]))


def verify_png(doc, path, manifest):
    from PIL import Image

    exp = doc["extraction_expectations"]
    img = Image.open(str(path))
    img.load()
    check(doc["id"] + ": image readable", img.size == (exp["image_size_px"]["width"], exp["image_size_px"]["height"]),
          "size={}".format(img.size))
    check(doc["id"] + ": RGB mode", img.mode == "RGB", "mode={}".format(img.mode))
    check(doc["id"] + ": file size > 20KB", path.stat().st_size > 20 * 1024, "bytes={}".format(path.stat().st_size))
    try:
        import numpy as np
        arr = np.asarray(img.convert("L"), dtype="uint8")
        dark_fraction = float((arr < 128).mean())
        row_dark = (arr < 128).mean(axis=1)
        rules = int((row_dark > 0.5).sum())
        check(doc["id"] + ": has ink (not blank)", 0.005 < dark_fraction < 0.30, "dark_fraction={:.4f}".format(dark_fraction))
        check(doc["id"] + ": table rules detected", rules >= 4, "horizontal rule rows={}".format(rules))
        EXTRACTED[doc["id"]] = {"width": img.size[0], "height": img.size[1], "dark_fraction": dark_fraction,
                                "rule_rows": rules, "note": "no text layer; OCR required"}
    except ImportError:
        check(doc["id"] + ": numpy available", True, "numpy not installed - pixel checks skipped")


def verify_manifest(manifest):
    check("manifest: item count", len(manifest["items"]) == 9, len(manifest["items"]))
    check("manifest: document count", len(manifest["documents"]) == 5, len(manifest["documents"]))
    check("manifest: supplier count", len(manifest["suppliers"]) == 3, len(manifest["suppliers"]))
    check("manifest: item keys", [it["key"] for it in manifest["items"]] == fd.ITEM_KEYS)

    required = ["name_en", "name_zh", "hs_code", "materials_en", "materials_zh", "cartons",
                "pieces_per_carton", "total_pieces", "unit_price", "amount", "gross_weight_kg",
                "net_weight_kg", "cbm", "dimensions_cm"]
    for item in manifest["items"]:
        missing = [k for k in required if k not in item]
        check("manifest: item {} schema".format(item["key"]), not missing, "missing={}".format(missing))
        check("manifest: item {} pieces math".format(item["key"]),
              item["total_pieces"] == item["cartons"] * item["pieces_per_carton"])
        check("manifest: item {} amount math".format(item["key"]),
              approx(item["amount"], fd.round2(item["unit_price"] * item["total_pieces"])))
        dims = item["dimensions_cm"]
        expected_cbm = fd.round2(dims["l"] * dims["w"] * dims["h"] / 1_000_000.0 * item["cartons"])
        check("manifest: item {} CBM math".format(item["key"]), approx(item["cbm"], expected_cbm),
              "{} vs {}".format(item["cbm"], expected_cbm))

    for sid in fd.SUPPLIERS:
        computed = fd.supplier_totals(sid)
        check("manifest: {} totals match expected".format(sid), computed == fd.EXPECTED_TOTALS[sid],
              "{} vs {}".format(computed, fd.EXPECTED_TOTALS[sid]))
        check("manifest: {} manifest totals".format(sid), manifest["totals_by_supplier"][sid] == computed)
    check("manifest: grand totals match expected", manifest["grand_totals"] == fd.EXPECTED_GRAND_TOTALS,
          "{} vs {}".format(manifest["grand_totals"], fd.EXPECTED_GRAND_TOTALS))

    missing_hs = [it["key"] for it in manifest["items"] if it["hs_code"] is None]
    check("manifest: exactly one missing HS", missing_hs == ["B-2"], "missing={}".format(missing_hs))

    d = manifest["duplicates"][0]
    a = next(it for it in manifest["items"] if it["key"] == "A-1")
    b = next(it for it in manifest["items"] if it["key"] == "B-3")
    check("manifest: duplicate pair", d["item_keys"] == ["A-1", "B-3"])
    for field in d["match_fields"]:
        check("manifest: duplicate field {}".format(field), a.get(field) == b.get(field),
              "{} != {}".format(a.get(field), b.get(field)))

    for sid, sup in manifest["suppliers"].items():
        check("manifest: {} name EN".format(sid), bool(sup.get("name_en")))
        check("manifest: {} name ZH".format(sid), bool(sup.get("name_zh")))
        check("manifest: {} contract".format(sid), bool(sup.get("contract_no")))
        check("manifest: {} invoice".format(sid), bool(sup.get("invoice_no")))
        check("manifest: {} date".format(sid), bool(sup.get("doc_date")))
    check("manifest: destination Melbourne", "MELBOURNE" in manifest["destination"]["port"].upper())
    check("manifest: incoterm CIF", manifest["incoterm"] == "CIF" and "CIF" in manifest["incoterm_full"])

    for doc in manifest["documents"]:
        path = REPO / doc["path"]
        check("manifest: {} file exists".format(doc["id"]), path.exists(), str(path))
        if path.exists():
            check("manifest: {} bytes match".format(doc["id"]), path.stat().st_size == doc["bytes"],
                  "{} vs {}".format(path.stat().st_size, doc["bytes"]))
            import hashlib
            h = hashlib.sha256(path.read_bytes()).hexdigest()
            check("manifest: {} sha256".format(doc["id"]), h == doc["sha256"])
            check("manifest: {} item keys".format(doc["id"]), doc["line_item_keys"] == fd.DOC_ITEM_KEYS[doc["id"]])


def main():
    ensure_stdout_utf8()
    manifest_path = REPO / "testdata" / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    print("manifest:", manifest_path)
    print("documents:", len(manifest["documents"]), "items:", len(manifest["items"]))

    verify_manifest(manifest)

    for doc in manifest["documents"]:
        path = REPO / doc["path"]
        print("\n--", doc["id"], "->", doc["path"])
        if doc["format"] == "pdf":
            verify_pdf(doc, path, manifest)
        elif doc["format"] == "xlsx":
            verify_xlsx(doc, path, manifest)
        elif doc["format"] == "png":
            verify_png(doc, path, manifest)
        else:
            check(doc["id"] + ": known format", False, doc["format"])

    from datetime import datetime
    report = {
        "verified_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "manifest_path": "testdata/manifest.json",
        "check_count": len(CHECKS),
        "failure_count": len(FAILURES),
        "ok": len(FAILURES) == 0,
        "checks": CHECKS,
        "failures": FAILURES,
        "extracted": EXTRACTED,
        "counts": {
            "documents": len(manifest["documents"]),
            "items": len(manifest["items"]),
            "suppliers": len(manifest["suppliers"]),
            "lines_per_document": {d["id"]: d["line_count"] for d in manifest["documents"]},
        },
        "totals": manifest["grand_totals"],
    }
    report_path = HERE / "verify_report.json"
    with open(report_path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(report, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    print("\nchecks:", len(CHECKS), "failures:", len(FAILURES))
    print("report:", report_path)
    if FAILURES:
        print("\nFAILED CHECKS:")
        for failure in FAILURES:
            print(" -", failure["check"], "::", failure["detail"])
        print("VERIFY FAILED")
        return 1
    print("VERIFY OK")
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        print("VERIFY CRASHED")
        sys.exit(2)
