# -*- coding: utf-8 -*-
"""Generate the heterogeneous supplier-document fixture set and manifest.json.

Usage (run with the bundled Python):
    C:\\Users\\Administrator\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe repo/testdata/_build/build_fixtures.py

Outputs (all under repo/testdata/):
    supplier_a_ningbo_homeline/packing_list.pdf        (reportlab, text layer)
    supplier_a_ningbo_homeline/commercial_invoice.pdf  (reportlab, text layer)
    supplier_b_foshan_sunrise/packing_list.xlsx        (openpyxl, reordered)
    supplier_b_foshan_sunrise/commercial_invoice.xlsx  (openpyxl, reordered)
    supplier_c_guangzhou_evergreen/packing_list_scan.png (Pillow, OCR-friendly)
    manifest.json
    _build/build_report.json
"""

import hashlib
import json
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
TESTDATA = HERE.parent
REPO = TESTDATA.parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

import fixture_data as fd  # noqa: E402

FONT_SIMHEI = Path("C:/Windows/Fonts/simhei.ttf")
FONT_MSYH = Path("C:/Windows/Fonts/msyh.ttc")
FONT_MSYH_BOLD = Path("C:/Windows/Fonts/msyhbd.ttc")

PDF_ENGINE = None  # "reportlab" or "pymupdf", decided at runtime


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------

def ensure_stdout_utf8():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def rel(path):
    return Path(path).resolve().relative_to(REPO).as_posix()


def now_iso():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def detect_pdf_engine():
    try:
        import reportlab  # noqa: F401
        from reportlab.pdfbase.ttfonts import TTFont  # noqa: F401
        from reportlab.platypus import SimpleDocTemplate  # noqa: F401
        return "reportlab"
    except Exception:
        import pymupdf  # noqa: F401
        return "pymupdf"


# ---------------------------------------------------------------------------
# PDF specification (shared by the reportlab and PyMuPDF renderers)
# ---------------------------------------------------------------------------

def _supplier_lines(sup):
    return [
        sup["name_en"],
        sup["address_en"],
        "TEL: {}   EMAIL: {}".format(sup["tel"], sup["email"]),
        sup["name_zh"],
        sup["address_zh"],
    ]


def _meta_lines(sup, doc_type):
    lines = [
        ("INVOICE NO. 发票号码", sup["invoice_no"]),
        ("CONTRACT NO. 合同号码", sup["contract_no"]),
        ("DATE 日期", sup["doc_date_text"]),
        ("INCOTERM 成交方式", sup["incoterm"]),
        ("PORT OF LOADING 起运港", sup["port_of_loading"]),
        ("DESTINATION 目的地", fd.DESTINATION["port"]),
        ("PAYMENT 付款方式", sup["payment"]),
    ]
    if doc_type == "commercial_invoice":
        lines.append(("ORIGIN 原产地", sup["origin"]))
    else:
        lines.append(("SHIPPING MARK 唛头", sup["shipping_mark"]))
    return lines


def _buyer_text(sup):
    return "\n".join([
        "BUYER / 收货人: " + sup["buyer_en"],
        sup["buyer_zh"],
        sup["buyer_address"],
    ])


def build_pdf_spec(sup, doc_type, items):
    totals = fd.supplier_totals(sup["id"])
    if doc_type == "packing_list":
        title_en, title_zh = "PACKING LIST", "装  箱  清  单"
        columns = [
            {"header": "NO.\n序号", "width": 30, "align": "center"},
            {"header": "DESCRIPTION\n品名", "width": 150, "align": "left"},
            {"header": "HS CODE\n海关编码", "width": 70, "align": "center"},
            {"header": "MATERIALS\n材质", "width": 100, "align": "left"},
            {"header": "CARTONS\n箱数", "width": 42, "align": "center"},
            {"header": "PCS/CTN\n每箱件数", "width": 42, "align": "center"},
            {"header": "TOTAL PCS\n总件数", "width": 48, "align": "center"},
            {"header": "L\n长(CM)", "width": 32, "align": "center"},
            {"header": "W\n宽(CM)", "width": 32, "align": "center"},
            {"header": "H\n高(CM)", "width": 32, "align": "center"},
            {"header": "CBM\n方数", "width": 38, "align": "center"},
            {"header": "G.W.\n毛重(KGS)", "width": 40, "align": "center"},
            {"header": "N.W.\n净重(KGS)", "width": 40, "align": "center"},
            {"header": "PACKING\n包装类型", "width": 46, "align": "center"},
        ]
        rows = []
        for it in items:
            rows.append([
                str(it["line_no"]),
                "{}\n{}".format(it["name_en"], it["name_zh"]),
                it["hs_code"] or "",
                "{}\n{}".format(it["material_en"], it["material_zh"]),
                str(it["cartons"]),
                str(it["pieces_per_carton"]),
                str(it["total_pieces"]),
                str(it["dimensions_cm"]["l"]),
                str(it["dimensions_cm"]["w"]),
                str(it["dimensions_cm"]["h"]),
                "{:.2f}".format(it["cbm"]),
                "{:,}".format(it["gross_weight_kg"]),
                "{:,}".format(it["net_weight_kg"]),
                "{}\n{}".format(it["packing_en"], it["packing_zh"]),
            ])
        total_row = [
            "TOTAL 合计", "", "", "",
            str(totals["cartons"]), "", str(totals["total_pieces"]), "", "", "",
            "{:.2f}".format(totals["cbm"]),
            "{:,}".format(totals["gross_weight_kg"]),
            "{:,}".format(totals["net_weight_kg"]),
            "",
        ]
        notes = [
            "PACKING: EXPORT CARTON / 纸箱包装",
            "SHIPPING MARK: N/M    ORIGIN: CHINA    INCOTERM: CIF MELBOURNE",
            "TOTAL: {} CARTONS / {} PCS / G.W. {:,} KGS / N.W. {:,} KGS / {:.2f} CBM".format(
                totals["cartons"], totals["total_pieces"], totals["gross_weight_kg"],
                totals["net_weight_kg"], totals["cbm"]),
            "WE HEREBY CERTIFY THAT THE GOODS ARE OF CHINESE ORIGIN.",
        ]
    else:
        title_en, title_zh = "COMMERCIAL INVOICE", "商  业  发  票"
        columns = [
            {"header": "NO.\n序号", "width": 30, "align": "center"},
            {"header": "DESCRIPTION\n品名", "width": 170, "align": "left"},
            {"header": "HS CODE\n海关编码", "width": 74, "align": "center"},
            {"header": "MATERIALS\n材质", "width": 110, "align": "left"},
            {"header": "CARTONS\n箱数", "width": 50, "align": "center"},
            {"header": "PCS/CTN\n每箱件数", "width": 50, "align": "center"},
            {"header": "TOTAL PCS\n计价件数", "width": 60, "align": "center"},
            {"header": "UNIT PRICE\n单价(USD)", "width": 62, "align": "center"},
            {"header": "AMOUNT\n金额(USD)", "width": 74, "align": "center"},
            {"header": "REMARKS\n备注", "width": 70, "align": "center"},
        ]
        rows = []
        for it in items:
            rows.append([
                str(it["line_no"]),
                "{}\n{}".format(it["name_en"], it["name_zh"]),
                it["hs_code"] or "",
                "{}\n{}".format(it["material_en"], it["material_zh"]),
                str(it["cartons"]),
                str(it["pieces_per_carton"]),
                str(it["total_pieces"]),
                "{:,.2f}".format(it["unit_price"]),
                "{:,.2f}".format(it["amount"]),
                "CIF MELBOURNE",
            ])
        total_row = [
            "TOTAL 合计", "", "", "",
            str(totals["cartons"]), "", str(totals["total_pieces"]), "",
            fd.money(totals["amount"]), "",
        ]
        notes = [
            "SAY TOTAL: US DOLLARS {} ONLY.".format(fd.amount_in_words(totals["amount"])),
            "TOTAL PACKING: {} CARTONS / {} PCS    G.W.: {:,} KGS    N.W.: {:,} KGS    CBM: {:.2f}".format(
                totals["cartons"], totals["total_pieces"], totals["gross_weight_kg"],
                totals["net_weight_kg"], totals["cbm"]),
            sup["bank"] or "",
            "PAYMENT: T/T 30% DEPOSIT, BALANCE AGAINST B/L COPY.",
            "WE HEREBY CERTIFY THAT THIS INVOICE IS TRUE AND CORRECT.",
        ]
    return {
        "doc_type": doc_type,
        "title_en": title_en,
        "title_zh": title_zh,
        "supplier_name_en": sup["name_en"],
        "supplier_lines": _supplier_lines(sup),
        "meta": _meta_lines(sup, doc_type),
        "buyer_text": _buyer_text(sup),
        "columns": columns,
        "rows": rows,
        "total_row": total_row,
        "notes": [n for n in notes if n],
        "footer": "PAGE 1 / 1",
    }


# ---------------------------------------------------------------------------
# PDF renderer: reportlab (preferred)
# ---------------------------------------------------------------------------

def _reportlab_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "FixtureSimHei" in pdfmetrics.getRegisteredFontNames():
        return "FixtureSimHei"
    if FONT_SIMHEI.exists():
        pdfmetrics.registerFont(TTFont("FixtureSimHei", str(FONT_SIMHEI)))
        return "FixtureSimHei"
    if FONT_MSYH.exists():
        pdfmetrics.registerFont(TTFont("FixtureSimHei", str(FONT_MSYH)))
        return "FixtureSimHei"
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light"


def render_pdf_reportlab(spec, path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    font = _reportlab_font()

    def esc(text):
        return escape(str(text)).replace("\n", "<br/>")

    def style(size, leading, align=0):
        return ParagraphStyle("s", fontName=font, fontSize=size, leading=leading,
                              alignment=align, spaceBefore=0, spaceAfter=0)

    title_style = style(19, 22.5, 1)
    sub_style = style(11.5, 14, 1)
    body = style(8.4, 10.4, 0)
    meta = style(8.2, 10.8, 0)
    head = style(7.2, 8.6, 1)
    cell_center = style(7.7, 9.2, 1)
    cell_left = style(7.7, 9.2, 0)
    total_style = style(7.7, 9.2, 1)

    doc = SimpleDocTemplate(
        str(path), pagesize=landscape(A4),
        leftMargin=24, rightMargin=24, topMargin=20, bottomMargin=20,
        title=spec["title_en"], author=spec["supplier_name_en"],
        subject=spec["doc_type"],
    )
    story = [
        Paragraph(esc(spec["title_en"]), title_style),
        Paragraph(esc(spec["title_zh"]), sub_style),
        Spacer(1, 5),
    ]

    left_cell = [Paragraph(esc(line), body) for line in spec["supplier_lines"]]
    right_cell = [Paragraph(esc("{}: {}".format(k, v)), meta) for k, v in spec["meta"]]
    outer = Table([[left_cell, right_cell]], colWidths=[450, 320])
    outer.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(outer)
    story.append(Spacer(1, 3))
    story.append(Paragraph(esc(spec["buyer_text"]), body))
    story.append(Spacer(1, 6))

    data = [[Paragraph(esc(c["header"]), head) for c in spec["columns"]]]
    for row in spec["rows"]:
        cells = []
        for value, column in zip(row, spec["columns"]):
            cells.append(Paragraph(esc(value), cell_left if column["align"] == "left" else cell_center))
        data.append(cells)
    data.append([Paragraph(esc(v), total_style) for v in spec["total_row"]])

    table = Table(data, colWidths=[c["width"] for c in spec["columns"]], repeatRows=1)
    last = len(data) - 1
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9E9E9")),
        ("BACKGROUND", (0, last), (-1, last), colors.HexColor("#F2F2F2")),
        ("SPAN", (0, last), (1, last)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(table)
    story.append(Spacer(1, 8))
    for note in spec["notes"]:
        story.append(Paragraph(esc(note), body))
    story.append(Spacer(1, 6))
    story.append(Paragraph(esc(spec["footer"]), meta))
    doc.build(story)


# ---------------------------------------------------------------------------
# PDF renderer: PyMuPDF fallback (same spec, same extractable content)
# ---------------------------------------------------------------------------

def render_pdf_pymupdf(spec, path):
    import pymupdf

    width, height = 842.0, 595.0
    doc = pymupdf.open()
    page = doc.new_page(width=width, height=height)
    fontname = "simhei"
    page.insert_font(fontname=fontname, fontfile=str(FONT_SIMHEI))
    font = pymupdf.Font(fontfile=str(FONT_SIMHEI))
    black = (0.0, 0.0, 0.0)

    def text_width(text, size):
        return font.text_length(str(text), fontsize=size)

    def draw(x, y, text, size=8.0, align="left"):
        text = str(text)
        w = text_width(text, size)
        if align == "center":
            x -= w / 2.0
        elif align == "right":
            x -= w
        page.insert_text((x, y), text, fontname=fontname, fontsize=size, color=black)

    def draw_multiline(x, y, text, size=7.6, leading=None, align="left"):
        leading = leading or size * 1.3
        for i, line in enumerate(str(text).split("\n")):
            draw(x, y + i * leading, line, size=size, align=align)

    draw(width / 2.0, 42, spec["title_en"], 20, "center")
    draw(width / 2.0, 62, spec["title_zh"], 12, "center")

    for i, line in enumerate(spec["supplier_lines"]):
        draw(24, 90 + i * 11.5, line, 8.2 if i == 0 else 7.2)

    for i, (label, value) in enumerate(spec["meta"]):
        draw(width - 24, 90 + i * 12, "{}: {}".format(label, value), 7.4, "right")

    block_bottom = 90 + max(len(spec["supplier_lines"]) * 11.5, len(spec["meta"]) * 12)
    buyer_y = block_bottom + 14
    draw_multiline(24, buyer_y, spec["buyer_text"], 7.6, 10.5)
    buyer_lines = spec["buyer_text"].count("\n") + 1
    table_top = buyer_y + buyer_lines * 10.5 + 18

    xs = [24.0]
    for column in spec["columns"]:
        xs.append(xs[-1] + column["width"])
    header_h = 30.0
    row_h = 42.0
    total_h = 22.0
    n_rows = len(spec["rows"])
    table_bottom = table_top + header_h + n_rows * row_h + total_h

    # header row background
    page.draw_rect(pymupdf.Rect(xs[0], table_top, xs[-1], table_top + header_h),
                   color=black, fill=(0.913, 0.913, 0.913), width=0.6)

    def cell_text(rect, text, size, align):
        lines = str(text).split("\n")
        line_h = size * 1.25
        total_text_h = len(lines) * line_h
        baseline = rect.y0 + (rect.height - total_text_h) / 2.0 + size
        for i, line in enumerate(lines):
            if align == "center":
                draw((rect.x0 + rect.x1) / 2.0, baseline + i * line_h, line, size, "center")
            else:
                draw(rect.x0 + 3, baseline + i * line_h, line, size, "left")

    # header + data rows
    for r in range(1 + n_rows):
        y0 = table_top + (header_h if r > 0 else 0) + (r - 1) * row_h if r > 0 else table_top
        h = header_h if r == 0 else row_h
        for c, column in enumerate(spec["columns"]):
            rect = pymupdf.Rect(xs[c], y0, xs[c + 1], y0 + h)
            page.draw_rect(rect, color=black, fill=None, width=0.6)
            if r == 0:
                cell_text(rect, column["header"], 6.7, "center")
            else:
                cell_text(rect, spec["rows"][r - 1][c], 7.2, column["align"])

    # total row: first two columns merged
    total_y = table_top + header_h + n_rows * row_h
    page.draw_rect(pymupdf.Rect(xs[0], total_y, xs[-1], total_y + total_h),
                   color=black, fill=(0.95, 0.95, 0.95), width=0.6)
    for c in range(2, len(spec["columns"])):
        page.draw_line((xs[c], total_y), (xs[c], total_y + total_h), color=black, width=0.6)
    merged_rect = pymupdf.Rect(xs[0], total_y, xs[2], total_y + total_h)
    cell_text(merged_rect, spec["total_row"][0], 7.2, "center")
    for c in range(2, len(spec["columns"])):
        rect = pymupdf.Rect(xs[c], total_y, xs[c + 1], total_y + total_h)
        cell_text(rect, spec["total_row"][c], 7.2, spec["columns"][c]["align"])

    note_y = table_bottom + 16
    for note in spec["notes"]:
        draw(24, note_y, note, 7.3)
        note_y += 11
    draw(24, height - 20, spec["footer"], 7.3)

    doc.save(str(path), garbage=3, deflate=True)
    doc.close()


def render_pdf(spec, path):
    global PDF_ENGINE
    if PDF_ENGINE == "reportlab":
        try:
            render_pdf_reportlab(spec, path)
            return "reportlab"
        except Exception as exc:  # fall back once, keep the build usable
            print("WARN: reportlab rendering failed ({}); falling back to PyMuPDF".format(exc))
            PDF_ENGINE = "pymupdf"
    render_pdf_pymupdf(spec, path)
    return "pymupdf"


# ---------------------------------------------------------------------------
# XLSX renderer (openpyxl)
# ---------------------------------------------------------------------------

def _xlsx_styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="000000")
    return {
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "header_fill": PatternFill("solid", fgColor="D9E1F2"),
        "total_fill": PatternFill("solid", fgColor="F2F2F2"),
        "font_base": Font(name="Microsoft YaHei", size=10),
        "font_small": Font(name="Microsoft YaHei", size=9),
        "font_bold": Font(name="Microsoft YaHei", size=10, bold=True),
        "font_title": Font(name="Microsoft YaHei", size=14, bold=True),
        "font_doc_title": Font(name="Microsoft YaHei", size=13, bold=True),
        "align_center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "align_left": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "align_right": Alignment(horizontal="right", vertical="center", wrap_text=True),
    }


def _apply_range(ws, cell_range, styles, font=None, fill=None, align=None, border=False, number_format=None):
    for row in ws[cell_range]:
        for cell in row:
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if align is not None:
                cell.alignment = align
            if border:
                cell.border = styles["border"]
            if number_format:
                cell.number_format = number_format


def build_packing_xlsx(sup, items, path):
    from openpyxl import Workbook
    styles = _xlsx_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Packing List"

    headers = [
        "NO.", "CARTONS\n箱数", "DESCRIPTION\n英文品名", "中文品名", "HS CODE\n海关编码",
        "MATERIAL\n材质", "PCS/CTN\n每箱件数", "TOTAL PCS\n总件数", "N.W.(KGS)\n净重",
        "G.W.(KGS)\n毛重", "MEAS.(CM)\nL*W*H", "CBM\n方数", "PACKING\n包装类型",
    ]
    widths = {"A": 7, "B": 9, "C": 30, "D": 18, "E": 16, "F": 30, "G": 10,
              "H": 11, "I": 12, "J": 12, "K": 17, "L": 8, "M": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # --- two title rows ---------------------------------------------------
    ws.merge_cells("A1:M1")
    ws["A1"] = sup["name_en"]
    ws.row_dimensions[1].height = 26
    _apply_range(ws, "A1:M1", styles, font=styles["font_title"],
                 align=styles["align_center"])

    ws.merge_cells("A2:H2")
    ws["A2"] = "PACKING LIST / 装箱清单"
    ws.merge_cells("I2:J2")
    ws["I2"] = "Invoice No.: {}".format(sup["invoice_no"])
    ws.merge_cells("K2:M2")
    ws["K2"] = "Date: {}".format(sup["doc_date"])
    ws.row_dimensions[2].height = 24
    _apply_range(ws, "A2:H2", styles, font=styles["font_doc_title"], align=styles["align_center"])
    _apply_range(ws, "I2:J2", styles, font=styles["font_bold"], align=styles["align_center"])
    _apply_range(ws, "K2:M2", styles, font=styles["font_bold"], align=styles["align_center"])

    # --- buyer / contract strip ------------------------------------------
    ws.merge_cells("A3:B3")
    ws["A3"] = "TO: " + sup["buyer_en"]
    ws.merge_cells("C3:H3")
    ws["C3"] = "{}  /  {}".format(sup["buyer_address"], sup["buyer_zh"])
    ws.merge_cells("I3:J3")
    ws["I3"] = "Contract No.: " + sup["contract_no"]
    ws.merge_cells("K3:M3")
    ws["K3"] = "Destination: " + fd.DESTINATION["port"]
    ws.row_dimensions[3].height = 30
    _apply_range(ws, "A3:B3", styles, font=styles["font_small"], align=styles["align_left"])
    _apply_range(ws, "C3:H3", styles, font=styles["font_small"], align=styles["align_left"])
    _apply_range(ws, "I3:J3", styles, font=styles["font_small"], align=styles["align_left"])
    _apply_range(ws, "K3:M3", styles, font=styles["font_small"], align=styles["align_left"])

    # --- header row -------------------------------------------------------
    for i, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=i, value=header)
        cell.font = styles["font_bold"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["align_center"]
        cell.border = styles["border"]
    ws.row_dimensions[4].height = 38

    # --- data rows --------------------------------------------------------
    for idx, item in enumerate(items):
        row = 5 + idx
        values = [
            item["line_no"], item["cartons"], item["name_en"], item["name_zh"],
            item["hs_code"], item["material_en"], item["pieces_per_carton"],
            item["total_pieces"], item["net_weight_kg"], item["gross_weight_kg"],
            item["dimensions_text"], item["cbm"],
            "{}/{}".format(item["packing_en"], item["packing_zh"]),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = styles["font_small"]
            cell.border = styles["border"]
            cell.alignment = styles["align_left"] if col in (3, 4, 6, 13) else styles["align_center"]
        for col, fmt in ((9, "#,##0"), (10, "#,##0"), (12, "#,##0.00")):
            ws.cell(row=row, column=col).number_format = fmt
        ws.row_dimensions[row].height = 34

    # --- totals row -------------------------------------------------------
    totals = fd.supplier_totals(sup["id"])
    ws["A8"] = "TOTAL"
    ws["B8"] = totals["cartons"]
    ws["H8"] = totals["total_pieces"]
    ws["I8"] = totals["net_weight_kg"]
    ws["J8"] = totals["gross_weight_kg"]
    ws["L8"] = totals["cbm"]
    _apply_range(ws, "A8:M8", styles, font=styles["font_bold"], fill=styles["total_fill"],
                 align=styles["align_center"], border=True)
    for cell in ("B8", "H8", "I8", "J8"):
        ws[cell].number_format = "#,##0"
    ws["L8"].number_format = "#,##0.00"
    ws.row_dimensions[8].height = 22

    # --- footer notes -----------------------------------------------------
    ws.merge_cells("A9:M9")
    ws["A9"] = ("All goods are packed in export cartons. 所有货物均为出口纸箱包装。  "
                "INCOTERM: CIF MELBOURNE   PAYMENT: T/T   ORIGIN: CHINA")
    ws.row_dimensions[9].height = 22
    _apply_range(ws, "A9:M9", styles, font=styles["font_small"], align=styles["align_left"])
    ws.merge_cells("A10:M10")
    ws["A10"] = ("SHIPPING MARK: N/M     TOTAL: {} CARTONS / {} PCS / G.W. {:,} KGS / "
                 "N.W. {:,} KGS / {:.2f} CBM").format(
        totals["cartons"], totals["total_pieces"], totals["gross_weight_kg"],
        totals["net_weight_kg"], totals["cbm"])
    ws.row_dimensions[10].height = 22
    _apply_range(ws, "A10:M10", styles, font=styles["font_small"], align=styles["align_left"])

    ws.freeze_panes = "A5"
    ws.print_area = "A1:M10"
    wb.save(str(path))


def build_invoice_xlsx(sup, items, path):
    from openpyxl import Workbook
    styles = _xlsx_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "Commercial Invoice"

    headers = [
        "NO.", "DESCRIPTION\n英文品名", "中文品名", "MODEL\n型号", "HS CODE\n海关编码",
        "MATERIAL\n材质", "CARTONS\n箱数", "PCS/CTN\n每箱件数", "TOTAL PCS\n计价件数",
        "UNIT PRICE(USD)\n单价", "AMOUNT(USD)\n金额",
    ]
    widths = {"A": 7, "B": 30, "C": 18, "D": 13, "E": 16, "F": 30, "G": 9,
              "H": 11, "I": 12, "J": 14, "K": 15}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # --- two title rows ---------------------------------------------------
    ws.merge_cells("A1:K1")
    ws["A1"] = sup["name_en"]
    ws.row_dimensions[1].height = 26
    _apply_range(ws, "A1:K1", styles, font=styles["font_title"], align=styles["align_center"])

    ws.merge_cells("A2:D2")
    ws["A2"] = "COMMERCIAL INVOICE / 商业发票"
    ws.merge_cells("E2:G2")
    ws["E2"] = "Invoice No.: {}".format(sup["invoice_no"])
    ws.merge_cells("H2:K2")
    ws["H2"] = "Date: {}".format(sup["doc_date"])
    ws.row_dimensions[2].height = 24
    _apply_range(ws, "A2:D2", styles, font=styles["font_doc_title"], align=styles["align_center"])
    _apply_range(ws, "E2:G2", styles, font=styles["font_bold"], align=styles["align_center"])
    _apply_range(ws, "H2:K2", styles, font=styles["font_bold"], align=styles["align_center"])

    # --- seller / buyer / contract strip ---------------------------------
    ws.merge_cells("A3:C3")
    ws["A3"] = "SELLER: " + sup["name_en"]
    ws.merge_cells("D3:G3")
    ws["D3"] = "CONTRACT NO.: " + sup["contract_no"]
    ws.merge_cells("H3:K3")
    ws["H3"] = "BUYER: " + sup["buyer_en"]
    ws.row_dimensions[3].height = 30
    _apply_range(ws, "A3:C3", styles, font=styles["font_small"], align=styles["align_left"])
    _apply_range(ws, "D3:G3", styles, font=styles["font_small"], align=styles["align_center"])
    _apply_range(ws, "H3:K3", styles, font=styles["font_small"], align=styles["align_left"])

    ws.merge_cells("A4:B4")
    ws["A4"] = "TERMS: CIF MELBOURNE"
    ws.merge_cells("C4:K4")
    ws["C4"] = ("PORT OF LOADING: {}   DESTINATION: {}   PAYMENT: {}   ORIGIN: {}").format(
        sup["port_of_loading"], fd.DESTINATION["port"], sup["payment"], sup["origin"])
    ws.row_dimensions[4].height = 24
    _apply_range(ws, "A4:B4", styles, font=styles["font_small"], align=styles["align_left"])
    _apply_range(ws, "C4:K4", styles, font=styles["font_small"], align=styles["align_center"])

    # --- header row -------------------------------------------------------
    for i, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=i, value=header)
        cell.font = styles["font_bold"]
        cell.fill = styles["header_fill"]
        cell.alignment = styles["align_center"]
        cell.border = styles["border"]
    ws.row_dimensions[5].height = 38

    # --- data rows --------------------------------------------------------
    for idx, item in enumerate(items):
        row = 6 + idx
        values = [
            item["line_no"], item["name_en"], item["name_zh"], item["model"],
            item["hs_code"], item["material_en"], item["cartons"],
            item["pieces_per_carton"], item["total_pieces"], item["unit_price"],
            item["amount"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = styles["font_small"]
            cell.border = styles["border"]
            cell.alignment = styles["align_left"] if col in (2, 3, 6) else styles["align_center"]
        ws.cell(row=row, column=10).number_format = "#,##0.00"
        ws.cell(row=row, column=11).number_format = "#,##0.00"
        ws.row_dimensions[row].height = 34

    # --- totals row -------------------------------------------------------
    totals = fd.supplier_totals(sup["id"])
    ws["A9"] = "TOTAL"
    ws["G9"] = totals["cartons"]
    ws["I9"] = totals["total_pieces"]
    ws["K9"] = totals["amount"]
    _apply_range(ws, "A9:K9", styles, font=styles["font_bold"], fill=styles["total_fill"],
                 align=styles["align_center"], border=True)
    ws["G9"].number_format = "#,##0"
    ws["I9"].number_format = "#,##0"
    ws["K9"].number_format = "#,##0.00"
    ws.row_dimensions[9].height = 22

    ws.merge_cells("A10:K10")
    ws["A10"] = "SAY TOTAL: US DOLLARS {} ONLY.".format(fd.amount_in_words(totals["amount"]))
    ws.row_dimensions[10].height = 22
    _apply_range(ws, "A10:K10", styles, font=styles["font_small"], align=styles["align_left"])

    ws.merge_cells("A11:K11")
    ws["A11"] = sup["bank"] or ""
    ws.row_dimensions[11].height = 20
    _apply_range(ws, "A11:K11", styles, font=styles["font_small"], align=styles["align_left"])

    ws.merge_cells("A12:K12")
    ws["A12"] = "AUTHORIZED SIGNATURE: ______________________     DATE: {}".format(sup["doc_date_text"])
    ws.row_dimensions[12].height = 34
    _apply_range(ws, "A12:K12", styles, font=styles["font_small"], align=styles["align_left"])

    ws.freeze_panes = "A6"
    ws.print_area = "A1:K12"
    wb.save(str(path))


# ---------------------------------------------------------------------------
# PNG renderer (Pillow) - Supplier C scanned-style packing list
# ---------------------------------------------------------------------------

def _pil_font(size, bold=False):
    from PIL import ImageFont
    candidates = [FONT_MSYH_BOLD if bold else FONT_MSYH, FONT_SIMHEI]
    for candidate in candidates:
        if candidate.exists():
            try:
                if candidate.suffix.lower() == ".ttc":
                    return ImageFont.truetype(str(candidate), size, index=0)
                return ImageFont.truetype(str(candidate), size)
            except Exception:
                continue
    raise RuntimeError("No usable CJK font found for Pillow rendering")


def build_packing_png(sup, items, path):
    from PIL import Image, ImageDraw

    width, height = 1800, 1150
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    black = (0, 0, 0)
    gray = (90, 90, 90)

    f_title = _pil_font(46, bold=True)
    f_sub = _pil_font(30)
    f_company = _pil_font(29, bold=True)
    f_company_cn = _pil_font(23)
    f_meta = _pil_font(20)
    f_header = _pil_font(20, bold=True)
    f_cell = _pil_font(20)
    f_cell_bold = _pil_font(20, bold=True)
    f_footer = _pil_font(19)

    def lines_centered(cx, cy, text, font, fill=black, spacing=1.28):
        lines = str(text).split("\n")
        line_h = font.size * spacing
        top = cy - (len(lines) * line_h) / 2.0 + line_h / 2.0
        for i, line in enumerate(lines):
            draw.text((cx, top + i * line_h), line, font=font, fill=fill, anchor="mm")

    def text_left(x, y, text, font, fill=black):
        draw.text((x, y), text, font=font, fill=fill, anchor="lm")

    # outer page border (scanned-page feel)
    draw.rectangle([36, 28, width - 36, height - 28], outline=gray, width=3)

    lines_centered(width / 2.0, 88, "PACKING LIST", f_title)
    lines_centered(width / 2.0, 136, "装  箱  清  单", f_sub)

    text_left(90, 190, sup["name_en"], f_company)
    text_left(90, 228, sup["name_zh"], f_company_cn)

    text_left(90, 285, "TO: " + sup["buyer_en"], f_meta)
    text_left(90, 315, sup["buyer_address"], f_meta)
    text_left(90, 345, "BUYER: " + sup["buyer_zh"], f_meta)

    right_x = 1090
    meta_lines = [
        "Invoice No.: " + sup["invoice_no"],
        "Contract No.: " + sup["contract_no"],
        "Date: " + sup["doc_date_text"],
        "Incoterm: " + sup["incoterm"],
        "Port of Loading: " + sup["port_of_loading"],
        "Destination: " + fd.DESTINATION["port"],
    ]
    for i, line in enumerate(meta_lines):
        text_left(right_x, 285 + i * 30, line, f_meta)

    draw.line([(70, 380), (width - 70, 380)], fill=gray, width=2)

    columns = [
        ("NO.\n序号", 80),
        ("DESCRIPTION\n品名", 500),
        ("HS CODE\n海关编码", 250),
        ("CARTONS\n箱数", 150),
        ("PCS/CTN\n每箱件数", 150),
        ("TOTAL PCS\n总件数", 170),
        ("G.W.(KGS)\n毛重", 170),
        ("CBM\n方数", 150),
    ]
    xs = [90]
    for _, w in columns:
        xs.append(xs[-1] + w)
    table_top = 405
    header_h = 78
    row_h = 100
    total_h = 62
    table_bottom = table_top + header_h + len(items) * row_h + total_h

    # header fill
    draw.rectangle([xs[0], table_top, xs[-1], table_top + header_h], fill=(232, 232, 232))
    # total-row fill
    total_y = table_bottom - total_h
    draw.rectangle([xs[0], total_y, xs[-1], table_bottom], fill=(242, 242, 242))

    # grid
    for x in xs:
        draw.line([(x, table_top), (x, table_bottom)], fill=black, width=2)
    y = table_top
    draw.line([(xs[0], y), (xs[-1], y)], fill=black, width=3)
    y += header_h
    draw.line([(xs[0], y), (xs[-1], y)], fill=black, width=2)
    for _ in items:
        y += row_h
        draw.line([(xs[0], y), (xs[-1], y)], fill=black, width=1)
    y += total_h
    draw.line([(xs[0], y), (xs[-1], y)], fill=black, width=3)

    # header labels
    header_cy = table_top + header_h / 2.0
    for i, (label, _) in enumerate(columns):
        lines_centered((xs[i] + xs[i + 1]) / 2.0, header_cy, label, f_header)

    # data rows
    for idx, item in enumerate(items):
        cy = table_top + header_h + row_h * idx + row_h / 2.0
        cells = [
            str(item["line_no"]),
            "{}\n{}".format(item["name_en"], item["name_zh"]),
            item["hs_code"] or "",
            str(item["cartons"]),
            str(item["pieces_per_carton"]),
            str(item["total_pieces"]),
            "{:,}".format(item["gross_weight_kg"]),
            "{:.2f}".format(item["cbm"]),
        ]
        for i, value in enumerate(cells):
            lines_centered((xs[i] + xs[i + 1]) / 2.0, cy, value, f_cell)

    # total row (first two columns merged)
    totals = fd.supplier_totals(sup["id"])
    total_cy = total_y + total_h / 2.0
    lines_centered((xs[0] + xs[2]) / 2.0, total_cy, "TOTAL 合计", f_cell_bold)
    total_values = {
        3: "{:,}".format(totals["cartons"]),
        5: "{:,}".format(totals["total_pieces"]),
        6: "{:,}".format(totals["gross_weight_kg"]),
        7: "{:.2f}".format(totals["cbm"]),
    }
    for i, value in total_values.items():
        lines_centered((xs[i] + xs[i + 1]) / 2.0, total_cy, value, f_cell_bold)

    # footer
    footer_y = table_bottom + 42
    draw.text((90, footer_y), "PACKING: EXPORT CARTON / 纸箱包装    SHIPPING MARK: N/M    ORIGIN: CHINA",
              font=f_footer, fill=black, anchor="lm")
    draw.text((90, footer_y + 36),
              "TOTAL: {} CARTONS / {} PCS / G.W. {:,} KGS / N.W. {:,} KGS / {:.2f} CBM".format(
                  totals["cartons"], totals["total_pieces"], totals["gross_weight_kg"],
                  totals["net_weight_kg"], totals["cbm"]),
              font=f_footer, fill=black, anchor="lm")
    draw.text((90, footer_y + 78), "SIGNED: ______________________      DATE: " + sup["doc_date_text"],
              font=f_footer, fill=black, anchor="lm")
    draw.text((width - 300, footer_y + 78), "PAGE 1 / 1", font=f_footer, fill=gray, anchor="lm")

    img.save(str(path), format="PNG", optimize=True)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    global PDF_ENGINE
    ensure_stdout_utf8()
    started = time.time()
    PDF_ENGINE = detect_pdf_engine()
    print("PDF engine:", PDF_ENGINE)

    for supplier in fd.SUPPLIERS.values():
        (REPO / fd.DOC_DIR[supplier["id"]]).mkdir(parents=True, exist_ok=True)

    files = []

    # --- Supplier A: reportlab PDFs ---
    sup_a = fd.SUPPLIERS["supplier_a"]
    items_a = fd.items_for("supplier_a")
    for doc_type, filename in (("packing_list", "packing_list.pdf"),
                               ("commercial_invoice", "commercial_invoice.pdf")):
        path = REPO / fd.DOC_DIR["supplier_a"] / filename
        spec = build_pdf_spec(sup_a, doc_type, items_a)
        engine = render_pdf(spec, path)
        files.append({"path": rel(path), "kind": "pdf", "engine": engine, "bytes": path.stat().st_size})
        print("built", rel(path), engine, path.stat().st_size, "bytes")

    # --- Supplier B: XLSX ---
    sup_b = fd.SUPPLIERS["supplier_b"]
    items_b = fd.items_for("supplier_b")
    packing_xlsx = REPO / fd.DOC_DIR["supplier_b"] / "packing_list.xlsx"
    invoice_xlsx = REPO / fd.DOC_DIR["supplier_b"] / "commercial_invoice.xlsx"
    build_packing_xlsx(sup_b, items_b, packing_xlsx)
    build_invoice_xlsx(sup_b, items_b, invoice_xlsx)
    for path in (packing_xlsx, invoice_xlsx):
        files.append({"path": rel(path), "kind": "xlsx", "engine": "openpyxl", "bytes": path.stat().st_size})
        print("built", rel(path), "openpyxl", path.stat().st_size, "bytes")

    # --- Supplier C: PNG ---
    sup_c = fd.SUPPLIERS["supplier_c"]
    items_c = fd.items_for("supplier_c")
    png_path = REPO / fd.DOC_DIR["supplier_c"] / "packing_list_scan.png"
    build_packing_png(sup_c, items_c, png_path)
    files.append({"path": rel(png_path), "kind": "png", "engine": "Pillow", "bytes": png_path.stat().st_size})
    print("built", rel(png_path), "Pillow", png_path.stat().st_size, "bytes")

    # --- document metadata + manifest ---
    documents = []
    for doc in fd.DOCUMENTS:
        path = REPO / fd.document_path(doc)
        if not path.exists():
            raise RuntimeError("missing generated document: {}".format(path))
        engine = "openpyxl" if doc["format"] == "xlsx" else ("Pillow" if doc["format"] == "png" else PDF_ENGINE)
        if doc["format"] == "pdf":
            used = [f["engine"] for f in files if f["path"] == rel(path)]
            engine = used[0] if used else PDF_ENGINE
        documents.append({
            "id": doc["id"],
            "supplier_id": doc["supplier_id"],
            "supplier_code": fd.SUPPLIERS[doc["supplier_id"]]["code"],
            "doc_type": doc["doc_type"],
            "doc_type_zh": doc["doc_type_zh"],
            "format": doc["format"],
            "path": fd.document_path(doc),
            "workspace_path": "repo/" + fd.document_path(doc),
            "filename": doc["filename"],
            "line_item_keys": fd.DOC_ITEM_KEYS[doc["id"]],
            "line_count": len(fd.DOC_ITEM_KEYS[doc["id"]]),
            "page_count": 1 if doc["format"] in ("pdf", "png") else None,
            "engine": engine,
            "mime_type": {"pdf": "application/pdf", "xlsx":
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                          "png": "image/png"}[doc["format"]],
            "bytes": path.stat().st_size,
            "sha256": sha256_file(path),
            "extraction_expectations": fd._extraction_expectations(doc),
        })

    manifest = fd.build_manifest(documents)
    manifest["created_at"] = now_iso()
    manifest["generator"] = {
        "script": "repo/testdata/_build/build_fixtures.py",
        "data_module": "repo/testdata/_build/fixture_data.py",
        "pdf_engine": PDF_ENGINE,
        "python": sys.version.split()[0],
    }
    manifest_path = REPO / "testdata" / "manifest.json"
    with open(manifest_path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(manifest, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    print("wrote", rel(manifest_path), manifest_path.stat().st_size, "bytes")

    report = {
        "generated_at": manifest["created_at"],
        "pdf_engine": PDF_ENGINE,
        "python": sys.version,
        "platform": sys.platform,
        "elapsed_seconds": round(time.time() - started, 2),
        "files": files,
        "counts": {
            "suppliers": len(fd.SUPPLIERS),
            "documents": len(documents),
            "items": len(fd.ITEMS),
            "lines_per_document": {doc["id"]: doc["line_count"] for doc in documents},
            "missing_hs_items": [it["key"] for it in fd.ITEMS if it["hs_code"] is None],
            "duplicate_pairs": [d["item_keys"] for d in manifest["duplicates"]],
        },
        "totals": {
            "by_supplier": manifest["totals_by_supplier"],
            "grand": manifest["grand_totals"],
        },
        "manifest_sha256": sha256_file(manifest_path),
        "manifest_bytes": manifest_path.stat().st_size,
    }
    report_path = HERE / "build_report.json"
    with open(report_path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(report, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    print("wrote", rel(report_path), report_path.stat().st_size, "bytes")

    print("SUMMARY: {} documents, {} items, grand amount USD {}, {} cartons, {:.2f} CBM".format(
        len(documents), len(fd.ITEMS), fd.money(manifest["grand_totals"]["amount"]),
        manifest["grand_totals"]["cartons"], manifest["grand_totals"]["cbm"]))
    print("BUILD OK")


if __name__ == "__main__":
    try:
        main()
    except Exception:
        traceback.print_exc()
        print("BUILD FAILED")
        sys.exit(1)
