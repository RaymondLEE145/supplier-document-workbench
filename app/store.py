# -*- coding: utf-8 -*-
"""Single-user persistent JSON state with a process-wide lock."""
from __future__ import annotations

import copy
import hashlib
import json
import shutil
import threading
from pathlib import Path

from openpyxl import load_workbook

from .config import DATA_DIR, DEFAULT_TEMPLATE, TEMPLATE_DIR, now_iso, new_id

STATE_PATH = DATA_DIR / "state.json"
_LOCK = threading.RLock()


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def inspect_template(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        sheets.append({
            "title": ws.title,
            "dim": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "merges": [str(r) for r in ws.merged_cells.ranges],
            "print_area": str(ws.print_area) if ws.print_area else "",
        })
    return {
        "filename": path.name,
        "path": str(path),
        "hash": file_sha256(path),
        "size": path.stat().st_size,
        "sheets": sheets,
        "sheet_names": [s["title"] for s in sheets],
        "uploaded_at": now_iso(),
    }


def _seed_state() -> dict:
    default_supplier = {
        "id": new_id("sup_"),
        "name": "供应商 1",
        "created_at": now_iso(),
        "note": "默认工作位",
        "color": "gold",
    }
    tpl_path = TEMPLATE_DIR / "template.xlsx"
    if not tpl_path.exists() and DEFAULT_TEMPLATE.exists():
        shutil.copyfile(DEFAULT_TEMPLATE, tpl_path)
    template = inspect_template(tpl_path) if tpl_path.exists() else None
    if template:
        template["filename"] = "出口文件(1).xlsx"
    return {
        "version": 1,
        "created_at": now_iso(),
        "suppliers": [default_supplier],
        "active_supplier_id": default_supplier["id"],
        "files": [],
        "documents": [],
        "jobs": {},
        "compliance": {
            "run_at": None,
            "records": [],
            "confirmed": {},
            "items": {},
            "summary": {},
            "attempts": [],
        },
        "template": template,
        "exports": [],
        "activity": [{"at": now_iso(), "kind": "system", "message": "工作台初始化完成，已载入 5-Sheet 目标模板"}],
        "settings": {},
    }


def load_state() -> dict:
    with _LOCK:
        if not STATE_PATH.exists():
            state = _seed_state()
            _write_state(state)
            return state
        try:
            state = json.loads(STATE_PATH.read_text(encoding="utf-8"))
        except Exception:
            backup = STATE_PATH.with_suffix(".broken.json")
            try:
                shutil.copyfile(STATE_PATH, backup)
            except Exception:
                pass
            state = _seed_state()
            _write_state(state)
            return state
        # migrations / defaults
        state.setdefault("suppliers", [])
        state.setdefault("active_supplier_id", state["suppliers"][0]["id"] if state["suppliers"] else None)
        state.setdefault("files", [])
        state.setdefault("documents", [])
        state.setdefault("jobs", {})
        state.setdefault("compliance", {"run_at": None, "records": [], "confirmed": {}, "items": {}, "summary": {}, "attempts": []})
        state.setdefault("exports", [])
        state.setdefault("activity", [])
        state.setdefault("settings", {})
        if not state.get("template"):
            tpl_path = TEMPLATE_DIR / "template.xlsx"
            if not tpl_path.exists() and DEFAULT_TEMPLATE.exists():
                shutil.copyfile(DEFAULT_TEMPLATE, tpl_path)
            if tpl_path.exists():
                state["template"] = inspect_template(tpl_path)
        return state


def _write_state(state: dict) -> None:
    tmp = STATE_PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(STATE_PATH)


def save_state(state: dict) -> None:
    with _LOCK:
        _write_state(state)


def mutate(fn):
    """Run fn(state) under the state lock and persist."""
    with _LOCK:
        state = load_state()
        result = fn(state)
        _write_state(state)
        return result


def snapshot() -> dict:
    with _LOCK:
        return copy.deepcopy(load_state())


def log_activity(state: dict, kind: str, message: str) -> None:
    state.setdefault("activity", []).insert(0, {"at": now_iso(), "kind": kind, "message": message})
    state["activity"] = state["activity"][:120]


def find_supplier(state: dict, supplier_id: str | None) -> dict | None:
    if not supplier_id:
        return None
    for s in state.get("suppliers", []):
        if s.get("id") == supplier_id:
            return s
    return None


def supplier_files(state: dict, supplier_id: str) -> list[dict]:
    return [f for f in state.get("files", []) if f.get("supplier_id") == supplier_id]


def supplier_documents(state: dict, supplier_id: str) -> list[dict]:
    file_ids = {f["id"] for f in supplier_files(state, supplier_id)}
    return [d for d in state.get("documents", []) if d.get("file_id") in file_ids]
