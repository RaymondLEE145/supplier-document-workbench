# -*- coding: utf-8 -*-
"""Independent end-to-end validation of the downloaded workbook.

Run: python tests/validate_download.py <downloaded.xlsx> [template.xlsx]
It does not use the server's own validator; it opens the downloaded file with
openpyxl and checks sheets, header style fingerprints, per-cell data, formulas,
merged cells and cached totals.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
file = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "Downloads" / "e2e_export_download.xlsx"
template = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "assets" / "template.xlsx"
state = json.loads((ROOT / "data" / "state.json").read_text(encoding="utf-8"))
sup = next((s for s in state["suppliers"] if s["name"] == "E2E验收供应商"), None)
payloads = sorted((ROOT / "data" / "scripts").glob("payload_*.json"), key=lambda p: p.stat().st_mtime, reverse=True)
if sup and sup["id"] in state.get("analyses", {}).get("by_supplier", {}):
    analysis = state["analyses"]["by_supplier"][sup["id"]]
    items = [dict(i) for i in analysis["items"]]
    # apply the user-confirmed official supplements exactly as the server does
    comp = state.get("compliance") or {}
    for rec in comp.get("records") or []:
        if rec.get("confirmed") and rec.get("status") == "supported" and rec.get("value"):
            for it in items:
                if it.get("id") == rec.get("item_id"):
                    if rec["field"] == "hs_code":
                        it["hs_code"] = rec["value"]
                        it["hs_verified"] = True
                    elif rec["field"] in {"origin_criterion", "preference_eligible", "product_name_en", "product_name_zh"}:
                        it[rec["field"]] = rec["value"]
    doc = analysis["doc"]
elif payloads:
    payload = json.loads(payloads[0].read_text(encoding="utf-8"))
    items = payload["items"]
    doc = payload["doc"]
    print("using exported payload", payloads[0].name)
else:
    raise AssertionError("no E2E supplier and no exported payload found")
expected = {
    "cartons": sum(int(i["cartons"]) for i in items),
    "pieces": sum(int(i["pieces"]) for i in items),
    "amount": round(sum(float(i["amount"]) for i in items), 2),
    "gross": round(sum(float(i["gross_weight"]) for i in items), 3),
    "cbm": round(sum(float(i["cbm"]) for i in items), 4),
}
print("items", len(items), "expected", expected)
assert len(items) == 5, f"expected 5 deduped items, got {len(items)}"

wb = load_workbook(file, data_only=False)
wbv = load_workbook(file, data_only=True)
tpl = load_workbook(template)
checks = []
def chk(name, ok, detail=""):
    checks.append((name, bool(ok), detail))
    print(("PASS" if ok else "FAIL"), name, detail)

chk("5 sheets", wb.sheetnames == tpl.sheetnames, str(wb.sheetnames))
def sh(name):
    return next(s for s in wb.worksheets if s.title == name)
def shv(name):
    return next(s for s in wbv.worksheets if s.title == name)
def tsh(name):
    return next(s for s in tpl.worksheets if s.title == name)

# header style fingerprints
style_diffs = []
for name in ["装箱清单", "发票", "申报要素", "销货确认书", "报关单 "]:
    ws, tws = sh(name), tsh(name)
    for coord in ["A1", "A4", "B4", "C4"] if name == "装箱清单" else ["A1"]:
        a, b = tws[coord], ws[coord]
        fa = (a.font.bold, a.font.sz, a.fill.fill_type, a.fill.fgColor.rgb if a.fill.fill_type else None,
              a.border.left.style, a.border.right.style, a.border.top.style, a.border.bottom.style, a.number_format)
        fb = (b.font.bold, b.font.sz, b.fill.fill_type, b.fill.fgColor.rgb if b.fill.fill_type else None,
              b.border.left.style, b.border.right.style, b.border.top.style, b.border.bottom.style, b.number_format)
        if fa != fb:
            style_diffs.append(f"{name}!{coord}")
chk("header styles preserved", not style_diffs, ",".join(style_diffs))

# packing
ws = sh("装箱清单")
for idx, it in enumerate(items):
    r = 6 + idx
    chk(f"packing C{r}", ws[f"C{r}"].value == (it.get("name_en") or it.get("name_zh")), str(ws[f"C{r}"].value))
    chk(f"packing D{r}", ws[f"D{r}"].value == (it.get("name_zh") or it.get("name_en")), str(ws[f"D{r}"].value))
    chk(f"packing E{r}", ws[f"E{r}"].value == it["hs_code"], f"{ws[f'E{r}'].value} vs {it['hs_code']}")
    chk(f"packing G{r}", int(ws[f"G{r}"].value or 0) == int(it["cartons"]), f"{ws[f'G{r}'].value} vs {it['cartons']}")
tr = 6 + len(items)
chk("packing L total formula", ws[f"L{tr}"].value == f"=SUM(L6:L{5+len(items)})", str(ws[f"L{tr}"].value))
chk("packing M total formula", ws[f"M{tr}"].value == f"=SUM(M6:M{5+len(items)})", str(ws[f"M{tr}"].value))
chk("packing cached CBM", abs(float(shv("装箱清单")[f"L{tr}"].value) - expected["cbm"]) < 0.02, f"{shv('装箱清单')[f'L{tr}'].value} vs {expected['cbm']}")
chk("packing cached GW", abs(float(shv("装箱清单")[f"M{tr}"].value) - expected["gross"]) < 0.02, f"{shv('装箱清单')[f'M{tr}'].value} vs {expected['gross']}")
chk("packing carton helper", int(ws[f"O{tr}"].value or 0) == expected["cartons"], f"{ws[f'O{tr}'].value}")
chk("packing piece helper", int(ws[f"P{tr}"].value or 0) == expected["pieces"], f"{ws[f'P{tr}'].value}")

# invoice
ws = sh("发票")
for idx, it in enumerate(items):
    r = 10 + idx
    chk(f"invoice B{r}", ws[f"B{r}"].value == (it.get("name_en") or it.get("name_zh")), str(ws[f"B{r}"].value))
    chk(f"invoice D{r}", ws[f"D{r}"].value == it["hs_code"], str(ws[f"D{r}"].value))
    chk(f"invoice J{r}", round(float(ws[f"J{r}"].value or 0), 2) == round(float(it["amount"]), 2), f"{ws[f'J{r}'].value}")
tr = 10 + len(items)
chk("invoice J total formula", ws[f"J{tr}"].value == f"=SUM(J10:J{9+len(items)})", str(ws[f"J{tr}"].value))
chk("invoice cached amount", abs(float(shv("发票")[f"J{tr}"].value) - expected["amount"]) < 0.02, f"{shv('发票')[f'J{tr}'].value} vs {expected['amount']}")

# sales
ws = sh("销货确认书")
for idx, it in enumerate(items):
    r = 21 + idx
    chk(f"sales C{r}", ws[f"C{r}"].value == (it.get("name_zh") or it.get("name_en")), str(ws[f"C{r}"].value))
    chk(f"sales F{r}", int(ws[f"F{r}"].value or 0) == int(it["pieces"]), f"{ws[f'F{r}'].value}")
    chk(f"sales H{r}", round(float(ws[f"H{r}"].value or 0), 2) == round(float(it["amount"]), 2), f"{ws[f'H{r}'].value}")
tr = 21 + len(items)
chk("sales total formula", ws[f"H{tr}"].value == f"=SUM(H21:H{20+len(items)})", str(ws[f"H{tr}"].value))
chk("sales cached amount", abs(float(shv("销货确认书")[f"H{tr}"].value) - expected["amount"]) < 0.02, f"{shv('销货确认书')[f'H{tr}'].value}")
chk("seller populated", bool(ws["C3"].value) and ws["C3"].value != "", str(ws["C3"].value))
chk("buyer address populated", bool(ws["C12"].value), str(ws["C12"].value))

# customs
ws = sh("报关单 ")
for idx, it in enumerate(items):
    r = 20 + idx * 2
    chk(f"customs B{r}", ws[f"B{r}"].value == it["hs_code"], f"{ws[f'B{r}'].value} vs {it['hs_code']}")
    chk(f"customs D{r} formula", ws[f"D{r}"].value == f"=销货确认书!C{21+idx}", str(ws[f"D{r}"].value))
    chk(f"customs I{r} formula", ws[f"I{r}"].value == f"=销货确认书!F{21+idx}", str(ws[f"I{r}"].value))
    chk(f"customs cached D{r}", ws[f"D{r}"].value.startswith("=") and shv("报关单 ")[f"D{r}"].value == (it.get("name_zh") or it.get("name_en")), str(shv("报关单 ")[f"D{r}"].value))
    chk(f"customs cached I{r}", ws[f"I{r}"].value.startswith("=") and int(shv("报关单 ")[f"I{r}"].value or 0) == int(it["pieces"]), str(shv("报关单 ")[f"I{r}"].value))
chk("customs contract", ws["A12"].value == doc["contract_no"], f"{ws['A12'].value} vs {doc['contract_no']}")
chk("customs trade country formula", ws["E12"].value == f"=销货确认书!F{26+len(items)-2}", str(ws["E12"].value))
chk("customs total cartons", int(ws["E14"].value or 0) == expected["cartons"], str(ws["E14"].value))
chk("customs total gross", abs(float(ws["F14"].value or 0) - expected["gross"]) < 0.02, str(ws["F14"].value))
chk("customs total net", bool(ws["H14"].value is not None), str(ws["H14"].value))

# formulas / merges / print setup
bad = []
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and any(x in c.value for x in ["#REF!", "#VALUE!", "#NAME?"]):
                bad.append(f"{s.title}!{c.coordinate}")
chk("no broken formulas", not bad, ",".join(bad))
chk("packing merges", len(sh("装箱清单").merged_cells.ranges) >= 4, str(len(sh("装箱清单").merged_cells.ranges)))
chk("customs merges", len(sh("报关单 ").merged_cells.ranges) >= 80, str(len(sh("报关单 ").merged_cells.ranges)))
ps = sh("装箱清单").sheet_properties.pageSetUpPr
chk("print fit one page", bool(ps and ps.fitToPage) or sh("装箱清单").page_setup.fitToWidth == 1,
    f"fitToPage={bool(ps and ps.fitToPage)} scale={sh('装箱清单').page_setup.scale}")

failed = [c for c in checks if not c[1]]
print(f"\nTOTAL {len(checks)} CHECKS, FAILED {len(failed)}")
print("OVERALL", "PASS" if not failed else "FAIL")
if failed:
    for name, _, detail in failed:
        print(" -", name, detail)
    sys.exit(1)
